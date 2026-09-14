# -*- coding: utf-8 -*-
u"""🚄 코레일 공식 「KTX 시간표」 XLSX → data/train_time.json

왜 이 방식인가 (2026-09-14)
  `/ja/daytrip/`(ソウルから日帰り・1泊2日)에 쓸 «소요시간»이 우리에게 없었다.
  운임 3,231쌍은 있는데 시간이 없어서 API 를 세 군데 두드렸다:
    🔴 네이버 지도 API — NCP Directions 5 공식 문서 「자동차에 한해서만 제공」. 애초에 불가
    🔴 ODsay        — 호출은 되는데 장거리에서 «틀린다»(대전 251분·천안 118분 전철 경로).
                      자세한 건 fetch-traveltime.js 머리말. 표본 3개로 판정할 뻔했다
    🟡 TAGO 열차정보 — 맞는 길이지만 활용신청이 필요하다(fetch-train-time.js)
  ⭐ 그런데 **코레일이 공식 XLSX 로 그냥 준다**. 신청도 키도 필요 없다.
     letskorail 「열차운임/시간표」 게시판 → 「KTX 시간표(YYYY. M. D. 기준)」

무엇을 만드나
  수도권 출발역(서울·용산·청량리·수서·행신·판교) → 전국 79개 역의 «가장 빠른 편» 소요시간.
  출발역별로 따로 남긴다 — 일본 여행자에게 「어느 역에서 타는가」는 숙소 위치에 따라 답이 다르다.
  (명동·홍대면 서울역/용산역, 강남이면 수서역이 가깝다.)

🔴 읽을 때 주의 — 실물을 열어 보고 알아낸 두 가지
  1. **`00:00` 은 「그 역에 서지 않는다」는 뜻**이다. 빈칸 대신 0시로 들어온다.
     이걸 시각으로 읽으면 소요시간이 통째로 망가진다.
  2. **한 헤더 행에 하행·상행이 좌우로 나란히** 있다. `비고` 열이 하행 블록의 끝이다.
     그 뒤에 다시 `열차번호`가 나오면 거기부터 상행이다 — 서울 출발이 아니므로 안 쓴다.

⚠️ 갱신: 시간표가 바뀌면 코레일에서 새 XLSX 를 받아 data/raw/ktx-timetable.xlsx 로 덮고 다시 돌린다.
   원본을 data/raw/ 에 남기는 이유는 «어느 시점 시간표인지»를 나중에 확인할 수 있어야 하기 때문이다.
   (data/ 는 .vercelignore 라 배포에는 올라가지 않는다.)

실행: py parse-train-timetable.py
"""
import openpyxl, datetime, json, os, sys

ROOT = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(ROOT, 'data', 'raw', 'ktx-timetable.xlsx')
SRC_FARE = os.path.join(ROOT, 'data', 'raw', 'ktx-fare.xls')   # 같은 게시판의 「KTX 운임표」
OUT = os.path.join(ROOT, 'data', 'train_time.json')

# 수도권 출발역만. ⚠️ 부전(부산)발 동해선이 섞이면 「서울에서 1시간」 같은 거짓말이 된다 — 실제로 겪었다.
ORIG = [u'서울', u'용산', u'청량리', u'수서', u'행신', u'판교(경기)']


def tmin(v):
    u"""셀 → 분. 00:00 은 «미정차»이므로 None."""
    if isinstance(v, datetime.time):
        m = v.hour * 60 + v.minute
        return None if m == 0 else m
    return None


def read_fare():
    u"""「KTX 운임표」(.xls) → {도착역: {출발역: 일반실 운임}}

    ⚠️ 시트가 28개이고 «노선별 경유지별»로 나뉜다(같은 서울~부산도 경주 경유·구포 경유가 따로다).
       그래서 같은 구간이 여러 시트에 나온다 → **가장 싼 값**을 택한다.
    ⚠️ 열 구조가 시트마다 다르다(10열·7열). 그래서 「일반실」이 있는 열을 «찾아서» 쓰고,
       그 왼쪽 두 칸을 구간(역A·역B)으로 본다. 열 번호를 외워서 쓰면 어느 시트에서 조용히 틀린다.
    """
    if not os.path.exists(SRC_FARE):
        print(u'⚠️ 운임표가 없다(%s) — 소요시간만 만든다' % SRC_FARE)
        return {}
    import xlrd
    wb = xlrd.open_workbook(SRC_FARE)
    pair = {}
    for sn in wb.sheet_names():
        sh = wb.sheet_by_name(sn)
        hi = fi = None
        for i in range(min(14, sh.nrows)):
            for j in range(sh.ncols):
                if u'일반실' in str(sh.cell_value(i, j)):
                    hi, fi = i, j
                    break
            if hi is not None:
                break
        if hi is None or fi < 2:
            continue
        a, b = fi - 2, fi - 1
        for i in range(hi + 1, sh.nrows):
            s1 = str(sh.cell_value(i, a)).strip()
            s2 = str(sh.cell_value(i, b)).strip()
            v = sh.cell_value(i, fi)
            if not s1 or not s2 or not isinstance(v, float) or v <= 0:
                continue
            for o, d in ((s1, s2), (s2, s1)):   # 운임은 방향이 같다
                if o in ORIG and d not in ORIG:
                    k = (o, d)
                    if k not in pair or v < pair[k]:
                        pair[k] = int(v)
    agg = {}
    for (o, d), v in pair.items():
        agg.setdefault(d, {})[o] = v
    return agg


def main():
    if not os.path.exists(SRC):
        raise SystemExit(u'✗ 원본이 없다: %s\n  코레일 「열차운임/시간표」에서 «KTX 시간표» XLSX 를 받아 이 경로에 둘 것' % SRC)
    wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
    st = {}
    names = {}     # 역 이름 → {han, en} — 시간표가 한자·영문을 같이 싣고 있다
    for sn in wb.sheetnames:
        rows = list(wb[sn].iter_rows(values_only=True))
        hi = hj = None
        for i, r in enumerate(rows):
            vals = [unicode(c).strip() if c is not None else u'' for c in r] if sys.version_info[0] == 2 \
                else [str(c).strip() if c is not None else '' for c in r]
            if u'열차번호' in vals:
                hi, hj = i, vals.index(u'열차번호')
                break
        if hi is None:
            continue
        hdr = [str(c).strip() if c is not None else '' for c in rows[hi]]
        # ⭐ 헤더 바로 아래 두 줄이 «한자»와 «영문» 역명이다. 일본어 페이지에 쓸 표기를 여기서 줍는다
        #    — 손으로 옮겨 적지 않는다. (한자는 중국어 간체가 섞여 있어 그대로는 못 쓴다: 首尔·东大邱)
        han = [str(c).strip() if c is not None else '' for c in rows[hi + 1]] if hi + 1 < len(rows) else []
        eng = [str(c).strip() if c is not None else '' for c in rows[hi + 2]] if hi + 2 < len(rows) else []
        start = hj + 2
        try:
            end = hdr.index(u'비고', start)      # ← 하행 블록의 끝
        except ValueError:
            end = len(hdr)
        cols = [(k, hdr[k]) for k in range(start, end) if hdr[k]]
        for k, nm in cols:
            e = names.setdefault(nm, {})
            if k < len(han) and han[k] and 'han' not in e:
                e['han'] = han[k]
            if k < len(eng) and eng[k] and 'en' not in e:
                e['en'] = eng[k]
        for r in rows[hi + 3:]:
            if not r or len(r) <= hj or r[hj] is None:
                continue
            kind = str(r[hj + 1]).strip() if len(r) > hj + 1 and r[hj + 1] else ''
            t = {}
            for k, nm in cols:
                if k < len(r):
                    m = tmin(r[k])
                    if m is not None:
                        t[nm] = m
            origs = [(o, t[o]) for o in ORIG if o in t]
            if not origs:
                continue
            for onm, ot in origs:
                for nm, at in t.items():
                    if nm in ORIG or at <= ot:
                        continue
                    d = at - ot
                    if d <= 0 or d > 600:        # 자정 넘김·이상치 방어
                        continue
                    e = st.setdefault(nm, {'byOrigin': {}, 'lines': set(), 'kinds': set(), 'trains': 0})
                    cur = e['byOrigin'].get(onm)
                    if cur is None or d < cur:
                        e['byOrigin'][onm] = d
                    e['lines'].add(sn)
                    e['kinds'].add(kind)
                    e['trains'] += 1
    wb.close()

    if len(st) < 30:
        raise SystemExit(u'✗ %d개 역밖에 못 읽었다 — 시트 구조가 바뀌었을 수 있다. 빈 파일을 만들지 않는다' % len(st))

    fares = read_fare()
    out = {}
    for nm, e in st.items():
        bo = e['byOrigin']
        best = min(bo, key=lambda k: bo[k])
        fa = fares.get(nm, {})
        rec = {
            'min': bo[best], 'from': best, 'byOrigin': bo,
            'trains': e['trains'],
            'lines': sorted(e['lines']),
            'kinds': sorted(x for x in e['kinds'] if x),
        }
        nz = names.get(nm, {})
        if nz.get('han'):
            rec['han'] = nz['han']
        if nz.get('en'):
            rec['en'] = nz['en']
        if fa:
            rec['fare'] = fa
            # «가장 빠른 출발역»의 운임을 대표값으로. 그 역 운임이 없으면 제일 싼 것.
            rec['fareBest'] = fa.get(best, min(fa.values()))
        out[nm] = rec
    mtime = datetime.date.fromtimestamp(os.path.getmtime(SRC)).isoformat()
    json.dump({
        'source': u'한국철도공사 「KTX 시간표」·「KTX 운임표」 (코레일 공식 배포)',
        'srcFile': 'data/raw/ktx-timetable.xlsx + ktx-fare.xls', 'srcDate': mtime,
        'note': u'所要時間は「いちばん速い列車」、運賃は一般室（普通車）。編成によってはもっとかかる。',
        'origins': ORIG,
        'updated': datetime.date.today().isoformat(),
        'stations': out,
    }, open(OUT, 'w', encoding='utf-8'), ensure_ascii=False)

    print(u'✓ data/train_time.json — %d개 역' % len(out))
    near = sorted(out, key=lambda x: out[x]['min'])
    for nm in near[:6]:
        print(u'   %-10s %3d분 (%s발)' % (nm, out[nm]['min'], out[nm]['from']))
    print(u'   …')
    for nm in near[-3:]:
        print(u'   %-10s %3d분 (%s발)' % (nm, out[nm]['min'], out[nm]['from']))


main()
